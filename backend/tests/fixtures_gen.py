"""
Generates small, real PDF/DOCX fixture files at test time (no binary blobs
committed to the repo). Used by tests/test_parsing_formats.py.
"""
import os
import tempfile


def make_pdf(path: str):
    import fitz
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((72, 72), "EMPLOYMENT AGREEMENT")
    p1.insert_text((72, 110), "1. TERMINATION")
    p1.insert_text((72, 130), "The Company may terminate without prior notice in cases of misconduct.")
    p2 = doc.new_page()
    p2.insert_text((72, 72), "2. COMPENSATION")
    p2.insert_text((72, 100), "The Company shall pay the Employee within 5 days of month end.")
    doc.save(path)
    doc.close()


def make_scanned_pdf(path: str):
    """A PDF page with no extractable text layer at all -> should trigger
    the OCR-fallback flagging path."""
    import fitz
    doc = fitz.open()
    doc.new_page()  # intentionally blank / no text layer
    doc.save(path)
    doc.close()


def make_docx(path: str):
    import docx
    d = docx.Document()
    d.add_paragraph("EMPLOYMENT AGREEMENT")
    d.add_paragraph("1. TERMINATION")
    d.add_paragraph("The Company may terminate without prior notice in cases of misconduct.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Party"
    table.rows[0].cells[1].text = "NovaTech Solutions Pvt. Ltd."
    d.save(path)


def tmp_file(suffix: str) -> str:
    fd, path = tempfile.mkstemp(suffix=suffix)
    os.close(fd)
    return path
