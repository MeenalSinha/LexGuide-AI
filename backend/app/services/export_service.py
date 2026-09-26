"""
MODULE 10 — export the consultation brief as TXT / DOCX / PDF.
Uses only libraries already in requirements.txt (python-docx, PyMuPDF) so no
new dependencies are introduced for this feature.
"""
import io


def brief_to_text(brief: dict) -> str:
    lines = [
        "LEGAL CONSULTATION BRIEF", "=" * 40, "",
        f"Matter: {brief['matter']}", "",
        brief["situation_summary"], "",
        "Documents Reviewed:",
    ]
    for d in brief["documents_reviewed"]:
        lines.append(f"  - {d['filename']} ({d['type']})")
    lines += ["", "Key Areas:"]
    for i, area in enumerate(brief["key_areas"], 1):
        lines.append(f"  {i}. {area}")
    lines += ["", "Questions:"]
    for i, q in enumerate(brief["questions"], 1):
        lines.append(f"  {i}. {q}")
    lines += ["", "Important Dates:"]
    for d in brief["important_dates"]:
        lines.append(f"  - {d['document']}: {d['date']} ({d['context']})")
    lines += ["", "Missing Information:"]
    for m in brief["missing_information"]:
        lines.append(f"  - {m}")
    lines += ["", "-" * 40, brief["disclaimer"]]
    return "\n".join(lines)


def brief_to_docx_bytes(brief: dict) -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("Legal Consultation Brief", level=1)
    d.add_paragraph(f"Matter: {brief['matter']}")
    d.add_paragraph(brief["situation_summary"])

    d.add_heading("Documents Reviewed", level=2)
    for doc in brief["documents_reviewed"]:
        d.add_paragraph(f"{doc['filename']} ({doc['type']})", style="List Bullet")

    d.add_heading("Key Areas", level=2)
    for area in brief["key_areas"]:
        d.add_paragraph(area, style="List Bullet")

    d.add_heading("Questions", level=2)
    for q in brief["questions"]:
        d.add_paragraph(q, style="List Number")

    d.add_heading("Important Dates", level=2)
    for date in brief["important_dates"]:
        d.add_paragraph(f"{date['document']}: {date['date']} ({date['context']})", style="List Bullet")

    d.add_heading("Missing Information", level=2)
    for m in brief["missing_information"]:
        d.add_paragraph(m, style="List Bullet")

    d.add_paragraph("")
    disclaimer = d.add_paragraph(brief["disclaimer"])
    disclaimer.italic = True

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def brief_to_pdf_bytes(brief: dict) -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(50, 50, 545, 792)
    text = brief_to_text(brief)
    page.insert_textbox(rect, text, fontsize=10, fontname="helv")
    buf = io.BytesIO(doc.tobytes())
    doc.close()
    return buf.getvalue()
